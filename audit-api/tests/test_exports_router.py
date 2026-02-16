from __future__ import annotations

import asyncio
from datetime import datetime, timezone
from io import BytesIO

from argon2 import PasswordHasher
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from app import auth, settings
from app.routers import auth as auth_router
from app.routers import exports
from tests.fakes import FakeConnection, FakePool, make_request


@pytest.fixture(autouse=True)
def _configure_auth(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(settings, "AUDIT_INGEST_API_KEY", "ingest-key")
    monkeypatch.setattr(settings, "AUDIT_READ_API_KEY", "read-key")
    monkeypatch.setattr(settings, "AUDIT_AUTH_USER", "audit-admin")
    monkeypatch.setattr(
        settings, "AUDIT_AUTH_PASSWORD_HASH", PasswordHasher().hash("super-secret")
    )
    monkeypatch.setattr(settings, "AUDIT_SESSION_SECRET", "session-secret")
    monkeypatch.setattr(settings, "AUDIT_SESSION_TTL_SECONDS", 1800)
    monkeypatch.setattr(settings, "AUDIT_ALLOW_UNAUTHENTICATED_READS", False)
    auth._failed_login_attempts.clear()


def _build_client(
    fetch_results: list[list[dict[str, object]]] | None = None,
) -> TestClient:
    app = FastAPI()
    app.state.db = FakePool(conn=FakeConnection(fetch_results=fetch_results or []))
    app.include_router(auth_router.router)
    app.include_router(exports.router)
    return TestClient(app)


def _parse_export_rows(payload: bytes) -> list[tuple[object, ...]]:
    workbook = load_workbook(filename=BytesIO(payload), read_only=True, data_only=True)
    sheet = workbook["Conversaties"]
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()
    return [tuple(row) for row in rows]


def test_export_rejects_unauthenticated_requests() -> None:
    with _build_client(fetch_results=[[]]) as client:
        response = client.get("/exports/logs.xlsx")

    assert response.status_code == 401


def test_export_accepts_session_cookie() -> None:
    with _build_client(fetch_results=[[]]) as client:
        login = client.post(
            "/auth/login",
            data={"username": "audit-admin", "password": "super-secret"},
            follow_redirects=False,
        )
        assert login.status_code == 303
        response = client.get("/exports/logs.xlsx")

    assert response.status_code == 200


def test_export_returns_headers_and_conversation_first_content() -> None:
    conversations_rows = [
        {
            "id": "session-123",
            "profile": "pixie-questioner",
            "started_at": datetime(2026, 2, 11, 9, 0, tzinfo=timezone.utc),
            "ended_at": datetime(2026, 2, 11, 9, 15, tzinfo=timezone.utc),
            "metadata": {"location": "Antwerpen"},
        }
    ]
    message_rows = [
        {
            "id": "msg-1",
            "conversation_id": "session-123",
            "role": "user",
            "content": "Hallo",
            "created_at": datetime(2026, 2, 11, 9, 1, tzinfo=timezone.utc),
        },
        {
            "id": "msg-2",
            "conversation_id": "session-123",
            "role": "assistant",
            "content": "Dag!",
            "created_at": datetime(2026, 2, 11, 9, 2, tzinfo=timezone.utc),
        },
        {
            "id": "msg-3",
            "conversation_id": "session-123",
            "role": "summary",
            "content": "Samenvatting van het gesprek",
            "created_at": datetime(2026, 2, 11, 9, 3, tzinfo=timezone.utc),
        },
    ]

    with _build_client(fetch_results=[conversations_rows, message_rows]) as client:
        response = client.get("/exports/logs.xlsx", headers={"X-API-Key": "read-key"})

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response.headers["content-disposition"].startswith(
        'attachment; filename="pixie-audit-export-'
    )

    rows = _parse_export_rows(response.content)
    assert rows[0] == (
        "Conversatie ID",
        "Profiel",
        "Locatie",
        "Gestart op",
        "Beeindigd op",
        "Aantal berichten",
        "Samenvatting",
        "Volledig gesprek",
    )
    assert rows[1] == (
        "session-123",
        "pixie-questioner",
        "Antwerpen",
        "2026-02-11 10:00:00",
        "2026-02-11 10:15:00",
        2,
        "Samenvatting van het gesprek",
        "[2026-02-11 10:01:00] user: Hallo\n[2026-02-11 10:02:00] assistant: Dag!",
    )


def test_export_uses_deterministic_query_ordering() -> None:
    conn = FakeConnection(
        fetch_results=[
            [
                {
                    "id": "session-123",
                    "profile": "pixie-questioner",
                    "started_at": datetime(2026, 2, 11, 9, 0, tzinfo=timezone.utc),
                    "ended_at": None,
                    "metadata": {"location": "Antwerpen"},
                }
            ],
            [],
        ]
    )
    request = make_request(FakePool(conn=conn))

    asyncio.run(exports.export_logs(request=request))

    assert len(conn.fetch_calls) == 2
    first_sql, first_params = conn.fetch_calls[0]
    second_sql, second_params = conn.fetch_calls[1]
    assert "order by started_at desc, id desc" in first_sql
    assert first_params == ()
    assert "where conversation_id = any($1::text[])" in second_sql
    assert "order by conversation_id asc, created_at asc, id asc" in second_sql
    assert second_params == (["session-123"],)


def test_export_selects_latest_summary_and_allows_missing_summary() -> None:
    conversations_rows = [
        {
            "id": "session-a",
            "profile": "pixie-questioner",
            "started_at": datetime(2026, 2, 11, 9, 0, tzinfo=timezone.utc),
            "ended_at": None,
            "metadata": {"location": "Antwerpen"},
        },
        {
            "id": "session-b",
            "profile": "pixie-questioner",
            "started_at": datetime(2026, 2, 10, 9, 0, tzinfo=timezone.utc),
            "ended_at": None,
            "metadata": {"location": "Gent"},
        },
    ]
    message_rows = [
        {
            "id": "001",
            "conversation_id": "session-a",
            "role": "summary",
            "content": "Oude samenvatting",
            "created_at": datetime(2026, 2, 11, 9, 4, tzinfo=timezone.utc),
        },
        {
            "id": "002",
            "conversation_id": "session-a",
            "role": "summary",
            "content": "Laatste samenvatting",
            "created_at": datetime(2026, 2, 11, 9, 5, tzinfo=timezone.utc),
        },
        {
            "id": "003",
            "conversation_id": "session-b",
            "role": "assistant",
            "content": "Geen summary hier",
            "created_at": datetime(2026, 2, 10, 9, 1, tzinfo=timezone.utc),
        },
    ]

    with _build_client(fetch_results=[conversations_rows, message_rows]) as client:
        response = client.get("/exports/logs.xlsx", headers={"X-API-Key": "read-key"})

    rows = _parse_export_rows(response.content)[1:]
    by_id = {str(row[0]): row for row in rows}

    assert by_id["session-a"][6] == "Laatste samenvatting"
    assert by_id["session-b"][6] in ("", None)
    assert by_id["session-b"][5] == 1


def test_export_formats_transcript_chronologically_and_escapes_formula_cells() -> None:
    conversations_rows = [
        {
            "id": "=session-danger",
            "profile": "+profile",
            "started_at": datetime(2026, 2, 11, 9, 0, tzinfo=timezone.utc),
            "ended_at": None,
            "metadata": {"location": " @locatie"},
        }
    ]
    message_rows = [
        {
            "id": "000",
            "conversation_id": "=session-danger",
            "role": "user",
            "content": "Eerste",
            "created_at": datetime(2026, 2, 11, 9, 1, tzinfo=timezone.utc),
        },
        {
            "id": "001",
            "conversation_id": "=session-danger",
            "role": "assistant",
            "content": "Tweede",
            "created_at": datetime(2026, 2, 11, 9, 2, tzinfo=timezone.utc),
        },
        {
            "id": "002",
            "conversation_id": "=session-danger",
            "role": "summary",
            "content": "-gevoelig",
            "created_at": datetime(2026, 2, 11, 9, 3, tzinfo=timezone.utc),
        },
    ]

    with _build_client(fetch_results=[conversations_rows, message_rows]) as client:
        response = client.get("/exports/logs.xlsx", headers={"X-API-Key": "read-key"})

    row = _parse_export_rows(response.content)[1]
    assert row[0] == "'=session-danger"
    assert row[1] == "'+profile"
    assert row[2] == "'@locatie"
    assert row[6] == "'-gevoelig"
    assert row[7] == (
        "[2026-02-11 10:01:00] user: Eerste\n[2026-02-11 10:02:00] assistant: Tweede"
    )
